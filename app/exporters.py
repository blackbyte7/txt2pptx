import html
import platform
import shutil
import subprocess
import tempfile
from pathlib import Path
from typing import Dict, List, Optional, Union

import pdf2image
import pptx

from log_config.logging_config import logger

# Conditional imports for Windows COM
if platform.system() == "Windows":
    try:
        import comtypes.client
        import pythoncom
        import win32com.client

        HAS_WIN_COM = True
    except ImportError:
        logger.warning(
            "Windows COM libraries (pywin32, comtypes) not found. PDF export on Windows will be disabled."
        )
        HAS_WIN_COM = False
else:
    HAS_WIN_COM = False


class PresentationExporter:
    def __init__(self, output_dir: Union[str, Path] = "exports"):
        self.output_dir = Path(output_dir)
        self.output_dir.mkdir(parents=True, exist_ok=True)
        self._check_dependencies()

    def _check_dependencies(self):
        """Checks for necessary external dependencies."""
        poppler_found = True
        pdf2image_lib_found = True

        if platform.system() != "Windows":
            unoconv_found = self._is_tool_installed("unoconv")
            soffice_found = self._is_tool_installed("soffice")
            if not unoconv_found and not soffice_found:
                logger.warning(
                    "Neither 'unoconv' nor 'soffice' (LibreOffice) found in PATH. PDF export on Unix-like systems may fail."
                )
        try:
            import pdf2image as _pdf2image

            pdf2image_lib_found = True
            try:
                dummy_path_for_check = "non_existent_dummy_file.pdf"
                _ = _pdf2image.pdfinfo_from_path(dummy_path_for_check)
                logger.debug(
                    "Poppler check: pdfinfo_from_path executed without PDFInfoNotInstalledError."
                )
            except _pdf2image.exceptions.PDFInfoNotInstalledError:
                logger.warning(
                    "'poppler' utility not found or not in PATH. Image export will be disabled. Install poppler for your OS."
                )
                poppler_found = False
            except FileNotFoundError:
                logger.debug(
                    "Poppler check: Passed (pdfinfo_from_path raised FileNotFoundError on dummy path)."
                )
            except Exception as poppler_check_err:
                logger.warning(f"Could not definitively check for poppler: {poppler_check_err}")
                poppler_found = False
        except ImportError:
            logger.warning(
                "'pdf2image' library not found. Image export will be disabled. Install with 'pip install pdf2image'."
            )
            pdf2image_lib_found = False

        self.can_export_images = pdf2image_lib_found and poppler_found
        self.can_export_pdf_unix = (
            (self._is_tool_installed("unoconv") or self._is_tool_installed("soffice"))
            if platform.system() != "Windows"
            else False
        )

    def _is_tool_installed(self, name: str) -> bool:
        """Check if a command-line tool is installed and in PATH."""
        try:
            subprocess.run(
                [name, "--version"],
                stdout=subprocess.PIPE,
                stderr=subprocess.PIPE,
                check=True,
            )
            return True
        except (FileNotFoundError, subprocess.CalledProcessError):
            return False

    def export_presentation(
        self,
        pptx_path: Union[str, Path],
        formats: Optional[List[str]] = None,
        output_name: Optional[str] = None,
    ) -> Dict[str, Union[str, List[str], Dict]]:
        """Exports the presentation to multiple formats."""
        if formats is None:
            formats = ["pdf", "html", "images"]

        pptx_path = Path(pptx_path)
        if not pptx_path.exists():
            logger.error(f"Input PPTX file not found: {pptx_path}")
            return {"error": f"Input PPTX file not found: {pptx_path}"}

        if not output_name:
            output_name = pptx_path.stem

        results: Dict[str, Union[str, List[str], Dict]] = {
            "original": str(pptx_path),
            "exports": {},
        }
        export_details: Dict[str, Union[str, List[str]]] = {}

        for format_type in formats:
            try:
                if format_type == "pdf":
                    pdf_path = self._export_to_pdf(pptx_path, output_name)
                    if pdf_path:
                        export_details["pdf"] = pdf_path
                elif format_type == "html":
                    html_path = self._export_to_html(pptx_path, output_name)
                    if html_path:
                        export_details["html"] = html_path
                elif format_type == "images":
                    pdf_for_images = export_details.get("pdf")
                    if not pdf_for_images or not isinstance(pdf_for_images, str):
                        logger.info("Generating temporary PDF for image export.")
                        pdf_for_images = self._export_to_pdf(
                            pptx_path, f"{output_name}_temp_for_images"
                        )

                    if pdf_for_images and isinstance(pdf_for_images, str):
                        image_paths = self._export_to_images(pdf_for_images, output_name)
                        if image_paths:
                            export_details["images"] = image_paths
                        if "temp_for_images" in str(pdf_for_images):
                            Path(pdf_for_images).unlink(missing_ok=True)
                    else:
                        raise RuntimeError("Could not generate PDF needed for image export.")

            except Exception as e:
                logger.error(f"Export to {format_type} failed: {e}", exc_info=True)
                export_details[format_type] = f"Export failed: {e}"

        results["exports"] = export_details
        return results

    def _export_to_pdf(self, pptx_path: Path, output_name: str) -> Optional[str]:
        """Exports PowerPoint to PDF using platform-specific methods."""
        output_path = self.output_dir / f"{output_name}.pdf"
        logger.info(f"Attempting PDF export: {pptx_path} -> {output_path}")

        try:
            if platform.system() == "Windows" and HAS_WIN_COM:
                return self._export_to_pdf_windows(pptx_path, output_path)
            elif platform.system() != "Windows":
                return self._export_to_pdf_unix(pptx_path, output_path)
            else:
                logger.warning("PDF export skipped: Unsupported platform or missing dependencies.")
                return None
        except Exception as e:
            logger.error(f"PDF export failed: {e}", exc_info=True)
            output_path.unlink(missing_ok=True)
            return None

    def _export_to_pdf_windows(self, pptx_path: Path, output_path: Path) -> Optional[str]:
        """Windows-specific PDF export using COM automation."""
        if not HAS_WIN_COM:
            logger.error("Cannot export PDF on Windows: pywin32 or comtypes not installed.")
            return None

        powerpoint = None
        deck = None
        co_initialized = False
        try:
            if "pythoncom" in globals():
                pythoncom.CoInitialize()
                co_initialized = True
        except Exception as co_err:
            logger.debug(f"COM CoInitialize note: {co_err}")

        try:
            abs_pptx_path = str(pptx_path.absolute())
            abs_output_path = str(output_path.absolute())
            logger.info("Using comtypes for PDF export...")
            powerpoint = comtypes.client.CreateObject("Powerpoint.Application")
            powerpoint.DisplayAlerts = 0

            deck = powerpoint.Presentations.Open(abs_pptx_path, WithWindow=False)
            deck.SaveAs(abs_output_path, 32)
            logger.info(f"Successfully exported PDF using comtypes: {output_path}")
            return str(output_path)

        except Exception as com_error:
            logger.warning(f"Comtypes PDF export failed: {com_error}. Trying win32com...")
            try:
                logger.info("Using win32com for PDF export...")
                powerpoint = win32com.client.Dispatch("Powerpoint.Application")
                powerpoint.DisplayAlerts = 0

                deck = powerpoint.Presentations.Open(abs_pptx_path, WithWindow=False)
                deck.SaveAs(abs_output_path, 32)  # 32 is ppSaveAsPDF
                logger.info(f"Successfully exported PDF using win32com: {output_path}")
                return str(output_path)

            except Exception as win32_error:
                logger.error(
                    f"Both comtypes and win32com PDF export methods failed on Windows: {win32_error}",
                    exc_info=True,
                )
                return None
        finally:
            try:
                if deck:
                    deck.Close()
                    logger.debug("Closed PowerPoint presentation object.")
            except Exception as close_err:
                logger.warning(f"Error closing PowerPoint presentation: {close_err}")
            try:
                if powerpoint:
                    powerpoint.Quit()
                    logger.debug("Quit PowerPoint application object.")
            except Exception as quit_err:
                logger.warning(f"Error quitting PowerPoint application: {quit_err}")
            if co_initialized and "pythoncom" in globals():
                try:
                    pythoncom.CoUninitialize()
                except Exception:
                    pass
            if platform.system() == "Windows" and HAS_WIN_COM:
                import gc

                gc.collect()

    def _export_to_pdf_unix(self, pptx_path: Path, output_path: Path) -> Optional[str]:
        """Unix-specific PDF export using LibreOffice/unoconv."""
        abs_pptx_path = str(pptx_path.absolute())

        try:
            if self._is_tool_installed("unoconv"):
                logger.info("Using unoconv for PDF export...")
                cmd = ["unoconv", "-f", "pdf", "-o", str(output_path), abs_pptx_path]
                process = subprocess.run(
                    cmd, check=True, capture_output=True, text=True, timeout=120
                )
                logger.debug(f"unoconv stdout: {process.stdout}")
                logger.debug(f"unoconv stderr: {process.stderr}")
                if output_path.exists():
                    logger.info(f"Successfully exported PDF using unoconv: {output_path}")
                    return str(output_path)
                else:
                    raise FileNotFoundError(
                        f"unoconv finished but output PDF not found: {output_path}"
                    )

            elif self._is_tool_installed("soffice"):
                logger.info("Using soffice (LibreOffice) for PDF export...")

                with tempfile.TemporaryDirectory() as temp_dir:
                    cmd = [
                        "soffice",
                        "--headless",
                        "--convert-to",
                        "pdf",
                        "--outdir",
                        temp_dir,
                        abs_pptx_path,
                    ]
                    process = subprocess.run(
                        cmd, check=True, capture_output=True, text=True, timeout=120
                    )
                    logger.debug(f"soffice stdout: {process.stdout}")
                    logger.debug(f"soffice stderr: {process.stderr}")

                    temp_pdf_path = Path(temp_dir) / f"{pptx_path.stem}.pdf"
                    if temp_pdf_path.exists():
                        output_path.unlink(missing_ok=True)
                        try:
                            temp_pdf_path.rename(output_path)
                        except OSError:
                            shutil.move(str(temp_pdf_path), str(output_path))
                        logger.info(f"Successfully exported PDF using soffice: {output_path}")
                        return str(output_path)
                    else:
                        raise FileNotFoundError(
                            f"soffice conversion likely failed, PDF not found in temp dir: {temp_pdf_path}"
                        )
            else:
                logger.error("Cannot export PDF on Unix: Neither unoconv nor soffice found.")
                return None

        except FileNotFoundError as fnf_err:
            logger.error(f"Tool not found for PDF export: {fnf_err}", exc_info=True)
            return None
        except subprocess.CalledProcessError as cpe:
            logger.error(f"PDF conversion command failed (exit code {cpe.returncode}): {cpe.cmd}")
            logger.error(f"stdout: {cpe.stdout}")
            logger.error(f"stderr: {cpe.stderr}")
            return None
        except subprocess.TimeoutExpired as timeout:
            logger.error(
                f"PDF conversion timed out after {timeout.timeout} seconds: {timeout.cmd}"
            )
            return None
        except Exception as e:
            logger.error(f"Unix PDF export failed: {e}", exc_info=True)
            return None

    def _export_to_html(self, pptx_path: Path, output_name: str) -> Optional[str]:
        """Basic HTML export with HTML escaping and table support."""
        output_path = self.output_dir / f"{output_name}.html"
        logger.info(f"Attempting basic HTML export (text only): {pptx_path} -> {output_path}")
        try:
            prs = pptx.Presentation(pptx_path)
            escaped_doc_title = html.escape(output_name)
            html_content = [
                "<!DOCTYPE html>",
                "<html>",
                "<head>",
                "<meta charset='utf-8'>",
                f"<title>{escaped_doc_title}</title>",
                "<style>",
                "body { font-family: sans-serif; margin: 20px; }",
                ".slide { border: 1px solid #ccc; padding: 15px; margin-bottom: 20px; page-break-after: always; }",
                ".slide h2 { color: #333; border-bottom: 1px solid #eee; padding-bottom: 5px; }",
                ".slide p { margin-left: 20px; }",
                ".slide table { border-collapse: collapse; margin: 10px 20px; width: 90%; }",
                ".slide th, .slide td { border: 1px solid #ccc; padding: 6px 10px; text-align: left; }",
                "</style>",
                "</head>",
                "<body>",
            ]
            for i, slide in enumerate(prs.slides):
                html_content.append(f"<div class='slide'><h2>Slide {i + 1}</h2>")
                title_shape = slide.shapes.title if hasattr(slide.shapes, "title") else None
                title = title_shape.text if title_shape else ""
                if title:
                    html_content.append(f"<h3>{html.escape(title)}</h3>")

                for shape in slide.shapes:
                    if title_shape is not None and shape is title_shape:
                        continue
                    if getattr(shape, "has_table", False) and hasattr(shape, "table"):
                        html_content.append("<table>")
                        for r_idx, row in enumerate(shape.table.rows):
                            tag = "th" if r_idx == 0 else "td"
                            cells_html = "".join(
                                f"<{tag}>{html.escape(cell.text)}</{tag}>" for cell in row.cells
                            )
                            html_content.append(f"<tr>{cells_html}</tr>")
                        html_content.append("</table>")
                    elif getattr(shape, "has_text_frame", False):
                        for paragraph in shape.text_frame.paragraphs:
                            if not paragraph.text:
                                continue
                            prefix = ""
                            if paragraph.level > 0:
                                prefix = "&nbsp;" * (paragraph.level * 4) + "- "
                            html_content.append(f"<p>{prefix}{html.escape(paragraph.text)}</p>")

                html_content.append("</div>")

            html_content.append("</body></html>")

            with open(output_path, "w", encoding="utf-8") as f:
                f.write("\n".join(html_content))

            logger.info(f"Successfully exported basic HTML: {output_path}")
            return str(output_path)
        except Exception as e:
            logger.error(f"Basic HTML export failed: {e}", exc_info=True)
            output_path.unlink(missing_ok=True)
            return None

    def _export_to_images(
        self, pdf_path: Union[str, Path], output_name: str
    ) -> Optional[List[str]]:
        """Exports presentation slides as images (requires PDF)."""
        pdf_path = Path(pdf_path)
        image_dir = self.output_dir / f"{output_name}_images"
        image_dir.mkdir(parents=True, exist_ok=True)
        logger.info(f"Attempting image export from PDF: {pdf_path} -> {image_dir}")

        try:
            images = pdf2image.convert_from_path(pdf_path, dpi=150)
        except ImportError:
            logger.error("Image export failed: 'pdf2image' library not installed.")
            return None
        except pdf2image.exceptions.PDFInfoNotInstalledError:
            logger.error("Image export failed: 'poppler' not found or not in PATH.")
            return None
        except Exception as e:
            logger.error(f"Failed to convert PDF to images: {e}", exc_info=True)
            return None

        image_paths = []
        try:
            for idx, image in enumerate(images):
                image_filename = f"slide_{idx + 1:03d}.png"
                image_path = image_dir / image_filename
                image.save(image_path, "PNG")
                image_paths.append(str(image_path))

            if not image_paths:
                logger.warning("Image export resulted in zero images.")
                return None

            logger.info(f"Successfully exported {len(image_paths)} images to {image_dir}")
            return image_paths
        except Exception as e:
            logger.error(f"Failed to save images: {e}", exc_info=True)
            for img_path in image_paths:
                Path(img_path).unlink(missing_ok=True)
            return None
