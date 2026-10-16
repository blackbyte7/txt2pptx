from typing import Any, Dict, Optional, Tuple

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_DATA_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import MSO_VERTICAL_ANCHOR, PP_PARAGRAPH_ALIGNMENT
from pptx.exc import PackageNotFoundError
from pptx.util import Inches, Pt

from log_config.logging_config import logger

# Default values
DEFAULT_FONT_NAME = "Calibri"
DEFAULT_FONT_SIZE = Pt(12)
DEFAULT_TITLE_FONT_SIZE = Pt(32)
DEFAULT_COLOR = RGBColor(0, 0, 0)
DEFAULT_SHAPE_POS = {
    "left": Inches(1.0),
    "top": Inches(1.5),
    "width": Inches(8.0),
    "height": Inches(5.0),
}


class PPTXGenerator:
    """Generates a PowerPoint presentation from a structured JSON."""

    def __init__(self, template: Optional[Dict] = None):
        """
        Initializes the generator with optional template data.

        Args:
            template (Optional[Dict]): A template dictionary containing colors, fonts,
                                        and layout mappings (e.g., from PresentationTemplates).
        """
        self.template = template or {}
        self.colors = self.template.get("colors", {})
        self.fonts = self.template.get("font", {})
        default_layouts = {
            "title_slide": 0,
            "title_content": 1,
            "section_header": 2,
            "two_content": 3,
            "comparison": 4,
            "title_only": 5,
            "blank": 6,
            "content_with_caption": 7,
            "picture_with_caption": 8,
            "opening": 0,
            "closing": 1,
            "content": 1,
            "table": 1,
            "chart": 1,
            "default": 1,
        }
        self.layout_map = {**default_layouts, **self.template.get("layouts", {})}
        logger.debug(
            f"PPTXGenerator initialized with template: {self.template.get('name', 'Default')}, "
            f"Layout Map Keys: {list(self.layout_map.keys())}"
        )

    def _get_color(self, role: str, default: RGBColor = DEFAULT_COLOR) -> RGBColor:
        """Gets a color from the template by role, with fallback."""
        hex_color = self.colors.get(role)
        if isinstance(hex_color, RGBColor):
            return hex_color
        if hex_color and isinstance(hex_color, str) and hex_color.startswith("#"):
            try:
                hex_val = hex_color[1:]
                if len(hex_val) == 6:
                    return RGBColor.from_string(hex_val)
                else:
                    logger.warning(
                        f"Invalid hex color format '{hex_color}' (must be #RRGGBB). Using default for role '{role}'."
                    )
                    return default
            except ValueError as ve:
                logger.warning(
                    f"Invalid hex value in color '{hex_color}' for role '{role}': {ve}. Using default."
                )
                return default
        elif not hex_color:
            return default
        else:
            logger.warning(f"Invalid color format '{hex_color}' for role '{role}'. Using default.")
            return default

    def _get_font_props(self, role: str, level: int = 0) -> Dict[str, Any]:
        """Gets font name and size from the template by role and level, supporting both str and dict configs."""
        raw_font_info = self.fonts.get(role, {})
        if isinstance(raw_font_info, str):
            font_info: Dict[str, Any] = {"name": raw_font_info}
        elif isinstance(raw_font_info, dict):
            font_info = raw_font_info
        else:
            logger.warning(
                f"Unexpected font configuration type {type(raw_font_info)} for role '{role}'. Using defaults."
            )
            font_info = {}

        base_font_size = DEFAULT_FONT_SIZE
        base_font_name = DEFAULT_FONT_NAME
        base_bold = False
        base_italic = False

        if role == "heading":
            base_font_size = DEFAULT_TITLE_FONT_SIZE
            base_bold = True

        template_font_name = font_info.get("name")
        template_size_pt_config = font_info.get("size")
        template_bold = font_info.get("bold")
        template_italic = font_info.get("italic")

        if template_font_name:
            base_font_name = template_font_name

        if template_bold is not None:
            base_bold = template_bold

        if template_italic is not None:
            base_italic = template_italic

        try:
            if isinstance(template_size_pt_config, (int, float)):
                parsed_size = Pt(int(template_size_pt_config))
                if parsed_size.pt > 0:
                    base_font_size = parsed_size
            elif isinstance(template_size_pt_config, str) and template_size_pt_config.isdigit():
                parsed_size = Pt(int(template_size_pt_config))
                if parsed_size.pt > 0:
                    base_font_size = parsed_size
            elif template_size_pt_config is not None:
                logger.warning(
                    f"Invalid font size format '{template_size_pt_config}' for role '{role}'. "
                    f"Using default/base {base_font_size.pt}pt."
                )
        except (ValueError, TypeError):
            logger.warning(
                f"Invalid font size value '{template_size_pt_config}' for role '{role}'. "
                f"Using default/base {base_font_size.pt}pt."
            )

        if level > 0:
            adjusted_size_pt = max(8, base_font_size.pt - level * 2)
            font_size = Pt(adjusted_size_pt)
            logger.debug(
                f"Adjusting font size for role '{role}', level {level}: {base_font_size.pt}pt -> {font_size.pt}pt"
            )
        else:
            font_size = base_font_size

        return {
            "name": base_font_name,
            "size": font_size,
            "bold": base_bold,
            "italic": base_italic,
        }

    def create_presentation_from_json(self, data: Dict[str, Any]) -> Optional[Presentation]:
        """Creates a pptx.Presentation object from the structured JSON data."""
        if not isinstance(data, dict) or "slides" not in data:
            logger.error("Invalid input data: Must be a dict with a 'slides' key.")
            return None

        slides_data = data.get("slides")
        if not isinstance(slides_data, list):
            logger.error(
                f"Invalid input data: 'slides' key must contain a list, but found type {type(slides_data)}."
            )
            return None

        try:
            prs = Presentation()
            logger.info(f"Starting presentation generation for {len(data['slides'])} slides.")

            if not prs.slide_layouts:
                logger.error(
                    "Failed to create presentation: The base presentation has no slide layouts."
                )
                return None

            for i, slide_data in enumerate(data.get("slides", [])):
                logger.debug(f"Generating Slide {i + 1}...")
                self._add_slide(prs, slide_data, i + 1)

            logger.info("Presentation generation completed.")
            return prs

        except PackageNotFoundError:
            logger.error(
                "Failed to create presentation. Ensure python-pptx is installed correctly."
            )
            return None
        except FileNotFoundError as fnf:
            logger.error(f"Failed to create presentation. Base template file not found: {fnf}")
            return None
        except Exception as e:
            logger.error(f"An error occurred during presentation generation: {e}", exc_info=True)
            return None

    def _add_slide(self, prs: Presentation, slide_data: Dict[str, Any], slide_number: int):
        """Adds a single slide to the presentation based on slide_data."""
        if not isinstance(slide_data, dict):
            logger.warning(
                f"Skipping slide {slide_number}: Invalid slide data format (not a dict)."
            )
            return

        slide_type = slide_data.get("slide_type", "content")
        design_elements = slide_data.get("design_elements") or {}
        layout_key_suggestion = (
            design_elements.get("layout", slide_type)
            if isinstance(design_elements, dict)
            else slide_type
        )

        layout_index = self.layout_map.get(
            layout_key_suggestion, self.layout_map.get("default", 1)
        )
        logger.debug(
            f"Slide {slide_number}: Type '{slide_type}', Layout Suggestion '{layout_key_suggestion}', "
            f"Mapped Index: {layout_index}"
        )

        try:
            if layout_index >= len(prs.slide_layouts):
                logger.warning(
                    f"Slide {slide_number}: Layout index {layout_index} (for key '{layout_key_suggestion}') "
                    f"is out of range for available layouts ({len(prs.slide_layouts)}). Falling back to layout 0."
                )
                layout_index = 0
            slide_layout = prs.slide_layouts[layout_index]
        except IndexError:
            logger.error(
                f"Slide {slide_number}: Cannot access any slide layouts (even index 0). Aborting slide creation."
            )
            return

        slide = prs.slides.add_slide(slide_layout)

        title_text = slide_data.get("title", "")
        if hasattr(slide.shapes, "title") and slide.shapes.title and title_text:
            try:
                title_shape = slide.shapes.title
                title_shape.text = str(title_text)
                if title_shape.has_text_frame and title_shape.text_frame.paragraphs:
                    p = title_shape.text_frame.paragraphs[0]
                    apply_template_font = True
                    try:
                        if (
                            p.font
                            and p.font.size
                            and isinstance(p.font.size, Pt)
                            and p.font.size.pt > 0
                        ):
                            apply_template_font = False
                            logger.debug(
                                f"Slide {slide_number}: Title placeholder seems to have existing font settings, retaining them."
                            )
                        elif p.font and p.font.size is not None:
                            logger.debug(
                                f"Slide {slide_number}: Title placeholder font size type {type(p.font.size)}, applying template font."
                            )
                            apply_template_font = True
                    except AttributeError:
                        logger.debug(
                            f"Slide {slide_number}: Could not determine existing title font size, applying template font."
                        )
                        apply_template_font = True

                    if apply_template_font:
                        font_props = self._get_font_props("heading")
                        title_color = self._get_color(
                            "accent1", self._get_color("text", DEFAULT_COLOR)
                        )
                        run = p.runs[0] if p.runs else p.add_run()
                        run.font.name = font_props["name"]
                        run.font.size = font_props.get("size", DEFAULT_TITLE_FONT_SIZE)
                        run.font.color.rgb = title_color
                        run.font.bold = font_props.get("bold", True)
                        run.font.italic = font_props.get("italic", False)
                else:
                    logger.warning(
                        f"Slide {slide_number}: Title shape found but lacks text frame/paragraph for formatting."
                    )
            except Exception as title_err:
                logger.warning(f"Slide {slide_number}: Failed to set or format title: {title_err}")
        elif title_text:
            logger.warning(
                f"Slide {slide_number}: Layout lacks title placeholder, adding title text box manually."
            )
            left, top, width, height = Inches(0.5), Inches(0.2), Inches(9.0), Inches(1.0)
            txBox = slide.shapes.add_textbox(left, top, width, height)
            tf = txBox.text_frame
            tf.alignment = PP_PARAGRAPH_ALIGNMENT.CENTER
            tf.word_wrap = True
            p = tf.paragraphs[0] if tf.paragraphs else tf.add_paragraph()
            p.clear()
            run = p.add_run()
            run.text = str(title_text)
            font_props = self._get_font_props("heading")
            title_color = self._get_color("accent1", self._get_color("text", DEFAULT_COLOR))
            run.font.name = font_props["name"]
            run.font.size = font_props.get("size", DEFAULT_TITLE_FONT_SIZE)
            run.font.color.rgb = title_color
            run.font.bold = font_props.get("bold", True)
            run.font.italic = font_props.get("italic", False)

        content_elements = slide_data.get("content_elements", [])
        if not isinstance(content_elements, list):
            logger.warning(
                f"Slide {slide_number}: 'content_elements' is not a list. Skipping content."
            )
            content_elements = []

        body_placeholder = None
        try:
            potential_indices = [1, 13, 14]
            for idx in potential_indices:
                if idx < len(slide.placeholders):
                    ph = slide.placeholders[idx]
                    if hasattr(ph, "is_content_placeholder") and ph.is_content_placeholder:
                        body_placeholder = ph
                        logger.debug(
                            f"Slide {slide_number}: Found body content placeholder at index {idx}."
                        )
                        break
            if not body_placeholder:
                for ph in slide.placeholders:
                    if (
                        hasattr(ph, "is_content_placeholder")
                        and ph.is_content_placeholder
                        and ph.placeholder_format.idx != 0
                    ):
                        body_placeholder = ph
                        logger.debug(
                            f"Slide {slide_number}: Found body placeholder via fallback logic at index {ph.placeholder_format.idx}."
                        )
                        break
        except (IndexError, KeyError, AttributeError) as ph_err:
            logger.warning(
                f"Slide {slide_number}: Error accessing placeholders: {ph_err}. Will add elements manually."
            )
            body_placeholder = None

        current_element_index_for_stacking = 0
        used_body_placeholder = False
        num_content_elements = len(content_elements)

        for i, element_data in enumerate(content_elements):
            if not isinstance(element_data, dict):
                logger.warning(
                    f"Slide {slide_number}: Skipping invalid content element {i + 1} (not a dict)."
                )
                continue

            element_type = element_data.get("element_type")
            position_data = element_data.get("position")

            left, top, width, height = self._get_element_position(
                position_data, i, num_content_elements
            )

            use_placeholder_for_this = (
                body_placeholder
                and element_type == "text"
                and not used_body_placeholder
                and num_content_elements == 1
            )

            try:
                if element_type == "text":
                    self._add_text_element(
                        slide,
                        element_data,
                        body_placeholder if use_placeholder_for_this else None,
                        left,
                        top,
                        width,
                        height,
                        slide_number,
                    )
                    if use_placeholder_for_this:
                        used_body_placeholder = True
                    if not use_placeholder_for_this:
                        current_element_index_for_stacking += 1

                elif element_type == "table":
                    self._add_table_element(
                        slide, element_data, left, top, width, height, slide_number
                    )
                    current_element_index_for_stacking += 1

                elif element_type == "chart":
                    self._add_chart_element(
                        slide, element_data, left, top, width, height, slide_number
                    )
                    current_element_index_for_stacking += 1

                elif element_type == "image_placeholder":
                    self._add_image_placeholder(
                        slide, element_data, left, top, width, height, slide_number
                    )
                    current_element_index_for_stacking += 1

                else:
                    logger.warning(
                        f"Slide {slide_number}: Skipping element {i + 1} with unknown type '{element_type}'."
                    )
                    current_element_index_for_stacking += 1

            except Exception as element_err:
                logger.error(
                    f"Slide {slide_number}: Failed to add element {i + 1} (type '{element_type}'): {element_err}",
                    exc_info=True,
                )
                try:
                    error_box = slide.shapes.add_textbox(left, top, width, Inches(0.5))
                    tf = error_box.text_frame
                    p = tf.add_paragraph()
                    p.text = f"Error adding element: {element_type}"
                    p.font.color.rgb = RGBColor(255, 0, 0)
                    p.font.size = Pt(10)
                except Exception as err_box_ex:
                    logger.error(
                        f"Slide {slide_number}: Could not even add error textbox: {err_box_ex}"
                    )
                current_element_index_for_stacking += 1

        notes_text = slide_data.get("speaker_notes")
        if notes_text and isinstance(notes_text, str):
            try:
                notes_slide = slide.notes_slide
                text_frame = notes_slide.notes_text_frame
                text_frame.clear()

                p = (
                    text_frame.paragraphs[0]
                    if text_frame.paragraphs
                    else text_frame.add_paragraph()
                )
                p.text = notes_text
                p.font.name = DEFAULT_FONT_NAME
                p.font.size = Pt(10)
                logger.debug(f"Slide {slide_number}: Added speaker notes.")
            except Exception as notes_err:
                logger.warning(f"Slide {slide_number}: Failed to add speaker notes: {notes_err}")

    def _get_element_position(
        self, position_data: Optional[Dict], element_index: int, total_elements: int
    ) -> Tuple[Inches, Inches, Inches, Inches]:
        """
        Calculates position based on JSON hint or defaults, with basic two-column layout.

        Args:
            position_data (Optional[Dict]): Explicit position hints from JSON.
            element_index (int): The 0-based index of the current element being placed.
            total_elements (int): The total number of content elements on the slide.

        Returns:
            Tuple[Inches, Inches, Inches, Inches]: left, top, width, height
        """
        logger.debug(
            f"Calculating position for element {element_index + 1} of {total_elements}."
        )

        if position_data and isinstance(position_data, dict):
            try:
                left_in = position_data.get("left_inches")
                top_in = position_data.get("top_inches")
                width_in = position_data.get("width_inches")
                height_in = position_data.get("height_inches")

                if all(
                    isinstance(v, (int, float))
                    for v in [left_in, top_in, width_in, height_in]
                ):
                    left = Inches(float(left_in))
                    top = Inches(float(top_in))
                    width = Inches(float(width_in))
                    height = Inches(float(height_in))
                    logger.debug(
                        f"Using position override: L={left.inches:.2f}, T={top.inches:.2f}, "
                        f"W={width.inches:.2f}, H={height.inches:.2f}"
                    )
                    return left, top, width, height
                else:
                    logger.warning(
                        f"Incomplete or invalid position override data {position_data}. Falling back to defaults."
                    )
            except (TypeError, ValueError) as pos_err:
                logger.warning(
                    f"Invalid position override data types {position_data}: {pos_err}. Falling back to defaults."
                )

        # Default Positioning Logic
        slide_width_inches = 10.0
        slide_height_inches = 7.5
        margin_left_inches = 0.5
        margin_right_inches = 0.5
        margin_top_inches = 1.5
        margin_bottom_inches = 0.5
        column_gap_inches = 0.4

        available_width = slide_width_inches - margin_left_inches - margin_right_inches
        available_height = slide_height_inches - margin_top_inches - margin_bottom_inches

        if total_elements <= 0:
            logger.warning(
                "Position calculation called with zero total elements. Returning minimal box."
            )
            return (
                Inches(margin_left_inches),
                Inches(margin_top_inches),
                Inches(1.0),
                Inches(1.0),
            )

        elif total_elements == 1:
            logger.debug("Applying single element full-width layout.")
            default_left = Inches(margin_left_inches)
            default_top = Inches(margin_top_inches)
            default_width = Inches(available_width)
            default_height = Inches(available_height)

        elif total_elements == 2:
            logger.debug("Applying basic two-column layout.")
            col_width = (available_width - column_gap_inches) / 2
            if col_width <= 0:
                col_width = available_width / 2
                column_gap_inches = 0

            if element_index == 0:
                default_left = Inches(margin_left_inches)
            else:
                default_left = Inches(margin_left_inches + col_width + column_gap_inches)

            default_top = Inches(margin_top_inches)
            default_width = Inches(col_width)
            default_height = Inches(available_height)

        else:
            logger.debug(f"Applying vertical stack layout (element {element_index + 1}).")
            gap_count = total_elements - 1
            total_gap_height = gap_count * 0.2
            element_height_inches = max(
                0.5, (available_height - total_gap_height) / total_elements
            )

            default_top = Inches(
                margin_top_inches + element_index * (element_height_inches + 0.2)
            )
            default_left = Inches(margin_left_inches)
            default_width = Inches(available_width)
            default_height = Inches(element_height_inches)

            if default_top.inches + default_height.inches > (
                slide_height_inches - margin_bottom_inches
            ):
                default_height = Inches(
                    max(
                        0.5,
                        (slide_height_inches - margin_bottom_inches) - default_top.inches,
                    )
                )
                if default_height.inches <= 0:
                    default_height = Inches(0.5)

        logger.debug(
            f"Calculated default position: L={default_left.inches:.2f}, T={default_top.inches:.2f}, "
            f"W={default_width.inches:.2f}, H={default_height.inches:.2f}"
        )
        return default_left, default_top, default_width, default_height

    def _add_text_element(
        self, slide, element_data, body_placeholder, left, top, width, height, slide_num
    ):
        """Adds a text element (bullets) to the slide, using placeholder if provided, else creates textbox."""
        bullets = element_data.get("bullet_points", [])
        text_title = element_data.get("title")
        level = element_data.get("level", 0) or 0

        if not bullets and not text_title:
            logger.debug(f"Slide {slide_num}: Skipping empty text element.")
            return

        logger.debug(
            f"Slide {slide_num}: Adding Text element (Title: {text_title is not None}, Bullets: {len(bullets)})"
        )

        tf = None
        if body_placeholder and body_placeholder.has_text_frame:
            tf = body_placeholder.text_frame
            tf.clear()
            logger.debug(
                f"Slide {slide_num}: Using provided body placeholder (idx {body_placeholder.placeholder_format.idx}) for text element."
            )
        else:
            logger.debug(f"Slide {slide_num}: Creating new textbox for text element.")
            txBox = slide.shapes.add_textbox(left, top, width, height)
            tf = txBox.text_frame

        if not tf:
            logger.error(f"Slide {slide_num}: Could not get or create text frame for text element.")
            return

        tf.word_wrap = True
        tf.vertical_anchor = MSO_VERTICAL_ANCHOR.TOP

        body_font_props = self._get_font_props("body", level=level)
        body_color = self._get_color("text")

        if text_title:
            p = tf.add_paragraph()
            p.text = str(text_title)
            p.level = max(0, level)
            run = p.runs[0] if p.runs else p.add_run()
            run.font.name = body_font_props["name"]
            run.font.size = Pt(body_font_props["size"].pt + 2)
            run.font.bold = True
            run.font.italic = body_font_props["italic"]
            run.font.color.rgb = body_color

        if bullets:
            para_offset = len(tf.paragraphs)
            for i, bullet_text in enumerate(bullets):
                if not isinstance(bullet_text, str):
                    logger.warning(
                        f"Slide {slide_num}, Text Element, Bullet {i + 1}: Invalid format (not a string). Skipping."
                    )
                    continue

                if (
                    i == 0
                    and para_offset == 0
                    and tf.paragraphs
                    and not tf.paragraphs[0].text
                ):
                    p = tf.paragraphs[0]
                else:
                    p = tf.add_paragraph()

                p.text = bullet_text
                p.level = max(0, level + 1)

                run = p.runs[0] if p.runs else p.add_run()
                run.font.name = body_font_props["name"]
                run.font.size = body_font_props["size"]
                run.font.bold = body_font_props["bold"]
                run.font.italic = body_font_props["italic"]
                run.font.color.rgb = body_color
        elif not text_title:
            logger.debug(f"Slide {slide_num}: Text element has no title or bullets.")

    def _add_table_element(
        self, slide, element_data, left, top, width, height, slide_num
    ):
        """Adds a table element to the slide."""
        headers = element_data.get("headers", [])
        rows_data = element_data.get("rows", [])
        table_title = element_data.get("title")

        if not headers or not rows_data:
            logger.warning(
                f"Slide {slide_num}: Skipping table element due to missing headers or rows."
            )
            return

        logger.debug(
            f"Slide {slide_num}: Adding Table element (Title: {table_title is not None}, "
            f"Rows: {len(rows_data)}, Cols: {len(headers)})"
        )

        num_rows = len(rows_data) + 1
        num_cols = len(headers)

        try:
            if table_title:
                title_height = Inches(0.4)
                if height < title_height:
                    logger.warning(
                        f"Slide {slide_num}: Table height too small for title. Adjusting."
                    )
                    title_height = height * 0.1
                title_box = slide.shapes.add_textbox(left, top, width, title_height)
                title_tf = title_box.text_frame
                p = title_tf.add_paragraph()
                p.text = str(table_title)
                font_props = self._get_font_props("body")
                run = p.runs[0] if p.runs else p.add_run()
                run.font.name = font_props["name"]
                run.font.size = Pt(font_props["size"].pt + 1)
                run.font.bold = True
                run.font.color.rgb = self._get_color("text")
                title_tf.vertical_anchor = MSO_VERTICAL_ANCHOR.BOTTOM
                top += title_height
                height -= title_height

            shape = slide.shapes.add_table(num_rows, num_cols, left, top, width, height)
            table = shape.table

            col_width = width / num_cols if num_cols > 0 else width
            for i in range(num_cols):
                table.columns[i].width = int(col_width)

            header_font_props = self._get_font_props("body")
            header_color = self._get_color("text")
            for j, header_text in enumerate(headers):
                cell = table.cell(0, j)
                cell.text = str(header_text)
                para = cell.text_frame.paragraphs[0]
                run = para.runs[0] if para.runs else para.add_run()
                run.font.name = header_font_props["name"]
                run.font.size = Pt(max(8, header_font_props["size"].pt - 1))
                run.font.bold = True
                run.font.color.rgb = header_color
                cell.vertical_anchor = MSO_VERTICAL_ANCHOR.MIDDLE
                para.alignment = PP_PARAGRAPH_ALIGNMENT.CENTER

            cell_font_props = self._get_font_props("body")
            cell_color = self._get_color("text")
            for i, row_data in enumerate(rows_data):
                for j, cell_text in enumerate(row_data):
                    if j >= num_cols:
                        continue
                    cell = table.cell(i + 1, j)
                    cell.text = str(cell_text)
                    para = cell.text_frame.paragraphs[0]
                    run = para.runs[0] if para.runs else para.add_run()
                    run.font.name = cell_font_props["name"]
                    run.font.size = cell_font_props["size"]
                    run.font.bold = cell_font_props["bold"]
                    run.font.italic = cell_font_props["italic"]
                    run.font.color.rgb = cell_color
                    cell.vertical_anchor = MSO_VERTICAL_ANCHOR.MIDDLE

        except Exception as table_err:
            logger.error(f"Slide {slide_num}: Failed to add table: {table_err}", exc_info=True)
            slide.shapes.add_textbox(
                left, top, width, Inches(0.5)
            ).text_frame.text = f"Error adding table: {table_err}"

    def _add_chart_element(
        self, slide, element_data, left, top, width, height, slide_num
    ):
        """Adds a chart element to the slide."""
        chart_type_str = element_data.get("chart_type")
        chart_title = element_data.get("title")
        categories = element_data.get("categories", [])
        series_data = element_data.get("series", {})

        if not chart_type_str or not categories or not series_data:
            logger.warning(
                f"Slide {slide_num}: Skipping chart element due to missing type, categories, or series."
            )
            return

        logger.debug(
            f"Slide {slide_num}: Adding Chart element (Type: {chart_type_str}, "
            f"Title: {chart_title is not None}, Categories: {len(categories)}, Series: {len(series_data)})"
        )

        try:
            chart_type_enum = getattr(XL_CHART_TYPE, str(chart_type_str).upper(), None)
            if chart_type_enum is None:
                logger.error(
                    f"Slide {slide_num}: Invalid chart type '{chart_type_str}'. Skipping chart."
                )
                slide.shapes.add_textbox(
                    left, top, width, Inches(0.5)
                ).text_frame.text = f"Invalid chart type: {chart_type_str}"
                return

            chart_data = CategoryChartData()
            chart_data.categories = categories

            for series_name, values in series_data.items():
                if len(values) != len(categories):
                    logger.warning(
                        f"Slide {slide_num}, Chart '{chart_title}': Series '{series_name}' data points ({len(values)}) "
                        f"mismatch category count ({len(categories)}). Check LLM output."
                    )
                    values = (list(values) + [0] * len(categories))[: len(categories)]
                try:
                    numeric_values = [
                        float(v)
                        if isinstance(v, (int, float, str))
                        and str(v).replace(".", "", 1).lstrip("-").isdigit()
                        else 0.0
                        for v in values
                    ]
                    chart_data.add_series(series_name, numeric_values)
                except (ValueError, TypeError) as series_err:
                    logger.error(
                        f"Slide {slide_num}, Chart '{chart_title}': Error processing series '{series_name}' "
                        f"data {values}: {series_err}. Skipping series."
                    )
                    continue

            if not chart_data.series:
                logger.error(
                    f"Slide {slide_num}, Chart '{chart_title}': No valid series data was added. Skipping chart."
                )
                return

            chart_shape = slide.shapes.add_chart(
                chart_type_enum, left, top, width, height, chart_data
            )
            chart = chart_shape.chart

            chart.has_legend = True
            chart.legend.position = XL_LEGEND_POSITION.BOTTOM
            chart.legend.include_in_layout = False

            if chart_title:
                chart.has_title = True
                chart.chart_title.text_frame.text = str(chart_title)
                if chart.chart_title.text_frame.paragraphs:
                    _ = (
                        chart.chart_title.text_frame.paragraphs[0].runs[0]
                        if chart.chart_title.text_frame.paragraphs[0].runs
                        else chart.chart_title.text_frame.paragraphs[0].add_run()
                    )

            try:
                plot = chart.plots[0]
                plot.has_data_labels = True
                data_labels = plot.data_labels
                data_labels.font.size = Pt(9)
                data_labels.font.color.rgb = self._get_color("text_light", RGBColor(80, 80, 80))
                data_labels.number_format = "0"
                data_labels.position = XL_DATA_LABEL_POSITION.OUTSIDE_END
            except AttributeError:
                logger.debug(
                    f"Slide {slide_num}, Chart '{chart_title}': Chart type might not support data labels configuration."
                )
            except IndexError:
                logger.warning(
                    f"Slide {slide_num}, Chart '{chart_title}': Could not access chart plot for data labels."
                )

        except Exception as chart_err:
            logger.error(f"Slide {slide_num}: Failed to add chart: {chart_err}", exc_info=True)
            slide.shapes.add_textbox(
                left, top, width, Inches(0.5)
            ).text_frame.text = f"Error adding chart: {chart_err}"

    def _add_image_placeholder(
        self, slide, element_data, left, top, width, height, slide_num
    ):
        """Adds a placeholder shape indicating where an image should go."""
        description = element_data.get("description", "Image placeholder")
        logger.debug(f"Slide {slide_num}: Adding Image Placeholder: '{description}'")
        try:
            placeholder_shape = slide.shapes.add_shape(
                MSO_AUTO_SHAPE_TYPE.RECTANGLE, left, top, width, height
            )
            tf = placeholder_shape.text_frame
            tf.text = f"[Image: {description}]"
            tf.word_wrap = True
            tf.vertical_anchor = MSO_VERTICAL_ANCHOR.MIDDLE
            p = tf.paragraphs[0]
            p.alignment = PP_PARAGRAPH_ALIGNMENT.CENTER
            run = p.runs[0] if p.runs else p.add_run()
            run.font.size = Pt(11)
            run.font.italic = True
            run.font.color.rgb = self._get_color("text_light", RGBColor(128, 128, 128))

            line = placeholder_shape.line
            line.color.rgb = self._get_color("accent2", RGBColor(180, 180, 180))
            line.dash_style = MSO_LINE_DASH_STYLE.DASH
            fill = placeholder_shape.fill
            fill.background()

        except Exception as img_placeholder_err:
            logger.error(
                f"Slide {slide_num}: Failed to add image placeholder: {img_placeholder_err}",
                exc_info=True,
            )