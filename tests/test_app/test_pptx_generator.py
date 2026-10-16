from unittest.mock import MagicMock, call, patch

import pytest
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.util import Inches, Pt

from app.pptx_generator import (
    DEFAULT_COLOR,
    DEFAULT_FONT_NAME,
    DEFAULT_FONT_SIZE,
    PPTXGenerator,
)


@pytest.fixture
def generator(sample_template_data):
    """ Creates a PPTXGenerator instance with sample template data. """
    return PPTXGenerator(template=sample_template_data)


@pytest.fixture
def generator_no_template():
    """ Creates a PPTXGenerator instance without a template. """
    return PPTXGenerator(template=None)


def test_generator_init_with_template(generator, sample_template_data):
    assert generator.template == sample_template_data
    assert generator.colors == sample_template_data["colors"]
    assert generator.fonts == sample_template_data["font"]
    assert "title_content" in generator.layout_map
    assert generator.layout_map["title_content"] == 1


def test_generator_init_without_template(generator_no_template):
    assert generator_no_template.template == {}
    assert generator_no_template.colors == {}
    assert generator_no_template.fonts == {}
    assert "title_content" in generator_no_template.layout_map
    assert generator_no_template.layout_map["title_content"] == 1


def test_get_color(generator):
    assert generator._get_color("text") == RGBColor(0x2C, 0x3E, 0x50)
    assert generator._get_color("accent1") == RGBColor(0x34, 0x98, 0xDB)
    assert generator._get_color("non_existent_role", default=RGBColor(1, 2, 3)) == RGBColor(1, 2, 3)
    assert generator._get_color("non_existent_role") == DEFAULT_COLOR
    generator.colors["invalid_hex"] = "#12345"
    assert generator._get_color("invalid_hex") == DEFAULT_COLOR
    generator.colors["invalid_chars"] = "#GGHHII"
    assert generator._get_color("invalid_chars") == DEFAULT_COLOR
    generator.colors["not_a_color"] = "blue"
    assert generator._get_color("not_a_color") == DEFAULT_COLOR


def test_get_font_props(generator):
    heading_props = generator._get_font_props("heading")
    assert heading_props["name"] == "Calibri"
    assert heading_props["size"] == Pt(32)
    assert heading_props["bold"] is True

    body_props = generator._get_font_props("body")
    assert body_props["name"] == "Calibri"
    assert body_props["size"] == Pt(12)
    assert body_props["bold"] is False

    body_level1_props = generator._get_font_props("body", level=1)
    assert body_level1_props["name"] == "Calibri"
    assert body_level1_props["size"] == Pt(10)

    unknown_props = generator._get_font_props("unknown_role")
    assert unknown_props["name"] == DEFAULT_FONT_NAME
    assert unknown_props["size"] == DEFAULT_FONT_SIZE

    generator.fonts["bad_size"] = {"name": "Arial", "size": "abc"}
    bad_size_props = generator._get_font_props("bad_size")
    assert bad_size_props["name"] == "Arial"
    assert bad_size_props["size"] == DEFAULT_FONT_SIZE


def test_get_font_props_with_string_font_config():
    """Verifies that templates specifying fonts as plain strings do not raise AttributeError."""
    str_font_template = {
        "name": "legacy_creative",
        "font": {"heading": "Helvetica Neue", "body": "Helvetica"},
    }
    gen = PPTXGenerator(template=str_font_template)
    heading_props = gen._get_font_props("heading")
    body_props = gen._get_font_props("body")

    assert heading_props["name"] == "Helvetica Neue"
    assert heading_props["size"] == Pt(32)
    assert heading_props["bold"] is True
    assert body_props["name"] == "Helvetica"
    assert body_props["size"] == DEFAULT_FONT_SIZE


@patch("app.pptx_generator.Presentation")
def test_create_presentation_from_json_success(
    mock_Presentation, generator, sample_detailed_presentation_json, mock_pptx_presentation
):
    mock_Presentation.return_value = mock_pptx_presentation

    json_input = sample_detailed_presentation_json.copy()
    json_input.pop("template_info", None)
    prs = generator.create_presentation_from_json(json_input)

    assert prs is mock_pptx_presentation
    mock_Presentation.assert_called_once()
    assert mock_pptx_presentation.slides.add_slide.call_count == len(
        sample_detailed_presentation_json["slides"]
    )


@patch("app.pptx_generator.Presentation")
def test_create_presentation_invalid_input(mock_Presentation, generator):
    assert generator.create_presentation_from_json(None) is None
    assert generator.create_presentation_from_json({"wrong_key": []}) is None
    assert generator.create_presentation_from_json({"slides": "not a list"}) is None
    mock_Presentation.assert_not_called()


@patch("app.pptx_generator.Presentation")
def test_create_presentation_no_layouts(
    mock_Presentation, generator, sample_detailed_presentation_json
):
    mock_pres = MagicMock()
    mock_pres.slide_layouts = []
    mock_Presentation.return_value = mock_pres

    json_input = sample_detailed_presentation_json.copy()
    json_input.pop("template_info", None)

    assert generator.create_presentation_from_json(json_input) is None
    mock_Presentation.assert_called_once()


def test_add_slide_layout_selection(generator, mock_pptx_presentation, mock_pptx_slide):
    """ Test selecting the correct layout index """
    mock_pptx_presentation.slides.add_slide.return_value = mock_pptx_slide
    mock_pptx_presentation.slide_layouts = [
        MagicMock(name="L0"),
        MagicMock(name="L1"),
        MagicMock(name="L2"),
    ]

    slide_data_1 = {
        "slide_type": "content",
        "title": "T1",
        "design_elements": {"layout": "title_content"},
    }
    generator._add_slide(mock_pptx_presentation, slide_data_1, 1)
    mock_pptx_presentation.slides.add_slide.assert_called_with(
        mock_pptx_presentation.slide_layouts[1]
    )

    slide_data_2 = {"slide_type": "section_header", "title": "T2"}
    generator._add_slide(mock_pptx_presentation, slide_data_2, 2)
    mock_pptx_presentation.slides.add_slide.assert_called_with(
        mock_pptx_presentation.slide_layouts[2]
    )

    slide_data_3 = {"slide_type": "unknown_type", "title": "T3"}
    generator._add_slide(mock_pptx_presentation, slide_data_3, 3)
    mock_pptx_presentation.slides.add_slide.assert_called_with(
        mock_pptx_presentation.slide_layouts[1]
    )

    generator.layout_map["out_of_range"] = 10
    slide_data_4 = {
        "slide_type": "content",
        "title": "T4",
        "design_elements": {"layout": "out_of_range"},
    }
    generator._add_slide(mock_pptx_presentation, slide_data_4, 4)
    mock_pptx_presentation.slides.add_slide.assert_called_with(
        mock_pptx_presentation.slide_layouts[0]
    )


def test_add_slide_title(
    generator, mock_pptx_presentation, mock_pptx_slide, sample_detailed_presentation_json
):
    mock_pptx_presentation.slides.add_slide.return_value = mock_pptx_slide
    slide_data = sample_detailed_presentation_json["slides"][0]
    mock_title_shape = mock_pptx_slide.shapes.title

    mock_run_for_title_ph = mock_title_shape.text_frame.paragraphs[0].add_run.return_value
    mock_run_for_title_ph.reset_mock()
    mock_title_shape.text_frame.paragraphs[0].add_run.reset_mock()
    mock_title_shape.text_frame.reset_mock()

    generator._add_slide(mock_pptx_presentation, slide_data, 1)

    assert mock_title_shape.text == slide_data["title"]
    assert mock_run_for_title_ph.font.size == Pt(32)

    mock_pptx_slide.shapes.title = None

    mock_textbox_tf_instance = mock_pptx_slide._mock_textbox_tf
    mock_textbox_tf_instance.reset_mock()
    mock_textbox_tf_instance.add_paragraph.reset_mock()

    mock_pptx_slide.shapes.add_textbox.reset_mock()

    generator._add_slide(mock_pptx_presentation, slide_data, 2)
    mock_pptx_slide.shapes.add_textbox.assert_called_once()

    assert mock_textbox_tf_instance.add_paragraph.called
    mock_p = mock_textbox_tf_instance.paragraphs[-1]

    assert mock_p.add_run.called
    mock_run_in_list = mock_p.runs[-1]

    assert mock_run_in_list.text == slide_data["title"]
    assert mock_run_in_list.font.size == Pt(32)


def test_add_slide_notes(
    generator, mock_pptx_presentation, mock_pptx_slide, sample_detailed_presentation_json
):
    mock_pptx_presentation.slides.add_slide.return_value = mock_pptx_slide
    slide_data = sample_detailed_presentation_json["slides"][0]

    generator._add_slide(mock_pptx_presentation, slide_data, 1)

    mock_notes_frame = mock_pptx_slide.notes_slide.notes_text_frame
    mock_notes_para = mock_notes_frame.add_paragraph.return_value
    assert mock_notes_para.text == slide_data["speaker_notes"]
    assert mock_notes_para.font.size == Pt(10)


def test_add_slide_handles_invalid_elements(generator, mock_pptx_presentation, mock_pptx_slide):
    mock_pptx_presentation.slides.add_slide.return_value = mock_pptx_slide
    slide_data = {
        "slide_number": 1,
        "slide_type": "content",
        "title": "Invalid Elements",
        "content_elements": [
            "not a dict",
            {"element_type": "unknown"},
            {"element_type": "text"},
        ],
    }
    generator._add_slide(mock_pptx_presentation, slide_data, 1)


def test_add_text_element(generator, mock_pptx_slide):
    text_data_bullets = {"element_type": "text", "bullet_points": ["Point 1", "Point 2"]}
    text_data_title_bullets = {
        "element_type": "text",
        "title": "Subsection",
        "bullet_points": ["Sub Point A"],
    }
    text_data_empty = {"element_type": "text"}

    mock_body_ph = mock_pptx_slide.placeholders[1]

    mock_body_ph.text_frame.reset_mock()
    generator._add_text_element(
        mock_pptx_slide,
        text_data_bullets,
        mock_body_ph,
        Inches(1),
        Inches(1.5),
        Inches(8),
        Inches(5),
        1,
    )
    mock_body_ph.text_frame.add_paragraph.assert_has_calls([call(), call()])
    paras = mock_body_ph.text_frame.add_paragraph.return_value
    assert paras.text == "Point 2"
    assert paras.level == 1

    mock_pptx_slide.shapes.add_textbox.reset_mock()
    generator._add_text_element(
        mock_pptx_slide,
        text_data_title_bullets,
        None,
        Inches(1),
        Inches(4),
        Inches(8),
        Inches(2),
        1,
    )
    mock_pptx_slide.shapes.add_textbox.assert_called_once()
    new_tf = mock_pptx_slide.shapes.add_textbox.return_value.text_frame
    assert new_tf.add_paragraph.call_count == 2

    mock_body_ph.text_frame.reset_mock()
    generator._add_text_element(
        mock_pptx_slide,
        text_data_empty,
        mock_body_ph,
        Inches(1),
        Inches(1.5),
        Inches(8),
        Inches(5),
        1,
    )
    mock_body_ph.text_frame.add_paragraph.assert_not_called()


@patch("app.pptx_generator.Inches", lambda x: x * 914400)
def test_add_table_element(generator, mock_pptx_slide, sample_detailed_presentation_json):
    table_data = sample_detailed_presentation_json["slides"][1]["content_elements"][0]
    mock_add_table = mock_pptx_slide.shapes.add_table
    mock_table_obj = mock_add_table.return_value.table
    mock_cell_func = mock_table_obj.cell
    mock_pptx_slide.shapes.add_textbox.reset_mock()

    generator._add_table_element(
        mock_pptx_slide, table_data, Inches(1), Inches(1.5), Inches(8), Inches(3), 2
    )

    mock_add_table.assert_called_once()

    mock_pptx_slide.shapes.add_textbox.assert_called_once()
    mock_cell_func.assert_any_call(0, 0)
    mock_cell_func.assert_any_call(1, 0)


@patch("app.pptx_generator.CategoryChartData")
@patch("app.pptx_generator.Inches", lambda x: x * 914400)
def test_add_chart_element(
    mock_ChartData, generator, mock_pptx_slide, sample_detailed_presentation_json
):
    chart_data_json = sample_detailed_presentation_json["slides"][2]["content_elements"][0]
    mock_add_chart = mock_pptx_slide.shapes.add_chart
    mock_add_chart.reset_mock()
    mock_chart_data_instance = mock_ChartData.return_value
    mock_chart_obj = mock_add_chart.return_value.chart

    generator._add_chart_element(
        mock_pptx_slide, chart_data_json, Inches(1), Inches(1.5), Inches(8), Inches(4), 3
    )

    mock_ChartData.assert_called_once()
    assert mock_chart_data_instance.categories == chart_data_json["categories"]
    mock_chart_data_instance.add_series.assert_has_calls(
        [
            call("Product A", [1500.0, 1200.0, 1800.0, 1350.0]),
            call("Product B", [800.0, 950.0, 700.0, 880.0]),
        ],
        any_order=True,
    )
    mock_add_chart.assert_called_once_with(
        XL_CHART_TYPE.BAR_CLUSTERED,
        Inches(1),
        Inches(1.5),
        Inches(8),
        Inches(4),
        mock_chart_data_instance,
    )

    assert mock_chart_obj.has_title is True
    assert mock_chart_obj.chart_title.text_frame.text == chart_data_json["title"]
    assert mock_chart_obj.has_legend is True
    assert mock_chart_obj.legend.position == XL_LEGEND_POSITION.BOTTOM


def test_add_image_placeholder_element(
    generator, mock_pptx_slide, sample_detailed_presentation_json
):
    img_data = sample_detailed_presentation_json["slides"][3]["content_elements"][1]
    mock_add_shape = mock_pptx_slide.shapes.add_shape
    mock_add_shape.reset_mock()
    mock_shape_obj = mock_add_shape.return_value

    generator._add_image_placeholder(
        mock_pptx_slide, img_data, Inches(7), Inches(5), Inches(2), Inches(1), 4
    )

    mock_add_shape.assert_called_once_with(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(7), Inches(5), Inches(2), Inches(1)
    )
    assert mock_shape_obj.text_frame.text == f"[Image: {img_data['description']}]"
    assert mock_shape_obj.line.dash_style == MSO_LINE_DASH_STYLE.DASH
