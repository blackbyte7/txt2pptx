import json
from unittest.mock import MagicMock

import pytest
from pptx.util import Pt


@pytest.fixture(scope="session")
def sample_style() -> str:
    return "Professional"


@pytest.fixture(scope="session")
def sample_template_name() -> str:
    return "professional"


@pytest.fixture(scope="session")
def sample_raw_text_content() -> str:
    return """
    Introduction to AI Testing
    Importance: Ensures reliability and accuracy of AI models. Challenges: Non-determinism, large state space.
    Types of Testing: Unit tests for components, Integration tests for workflow, Performance tests for speed, Robustness tests for edge cases.
    Sales Data Q1 2024
    Region, Product A, Product B
    North, 1500, 800
    South, 1200, 950
    East, 1800, 700
    West, 1350, 880
    Chart Idea: Show Product A vs Product B sales per region using a clustered bar chart.
    Conclusion: AI testing is vital. Future work involves adversarial testing. Contact: test@example.com
    """


@pytest.fixture(scope="session")
def sample_template_data() -> dict:
    return {
        "name": "professional",
        "description": "Clean, modern, corporate style.",
        "colors": {
            "background": "#FFFFFF",
            "text": "#2C3E50",
            "accent1": "#3498DB",
            "accent2": "#2980B9",
            "text_light": "#888888",
        },
        "font": {
            "heading": {"name": "Calibri", "size": 32, "bold": True},
            "body": {"name": "Calibri", "size": 12}
        },
        "layouts": {
            "title_slide": 0, "title_content": 1, "section_header": 2,
            "two_content": 3, "comparison": 4, "title_only": 5, "blank": 6,
            "content_with_caption": 7, "picture_with_caption": 8,
            "opening": 0, "content": 1, "closing": 1, "table": 1, "chart": 1, "default": 1,
        }
    }


@pytest.fixture(scope="session")
def sample_detailed_presentation_json(sample_template_data) -> dict:
    """ A more detailed JSON structure matching the updated tools output """
    return {
        "metadata": {
            "estimated_duration_minutes": 5,
            "target_audience": "Technical Team",
            "complexity_level": "intermediate",
            "presentation_style": "Professional",
            "template_name": "professional"
        },
        "slides": [
            {
                "slide_number": 1,
                "slide_type": "opening",
                "title": "Introduction to AI Testing",
                "content_elements": [
                    {
                        "element_type": "text",
                        "bullet_points": [
                            "Ensuring reliability and accuracy of AI models",
                            "Overcoming challenges like non-determinism"
                        ]
                    }
                ],
                "design_elements": {"layout": "title_slide"},
                "speaker_notes": "Start with the importance of the topic."
            },
            {
                "slide_number": 2,
                "slide_type": "content",
                "title": "Sales Data Q1 2024",
                "content_elements": [
                    {
                        "element_type": "table",
                        "title": "Regional Sales Performance",
                        "headers": ["Region", "Product A", "Product B"],
                        "rows": [
                            ["North", 1500, 800],
                            ["South", 1200, 950],
                            ["East", 1800, 700],
                            ["West", 1350, 880]
                        ],
                        "position": {
                            "left_inches": 1.0, "top_inches": 1.5,
                            "width_inches": 8.0, "height_inches": 3.0
                        }
                    }
                ],
                "design_elements": {"layout": "title_content"},
                "speaker_notes": "Highlight East region's performance for Product A."
            },
            {
                "slide_number": 3,
                "slide_type": "content",
                "title": "Sales Comparison",
                "content_elements": [
                    {
                        "element_type": "chart",
                        "chart_type": "BAR_CLUSTERED",
                        "title": "Product A vs Product B Sales by Region",
                        "categories": ["North", "South", "East", "West"],
                        "series": {
                            "Product A": [1500, 1200, 1800, 1350],
                            "Product B": [800, 950, 700, 880]
                        }
                    }
                ],
                "design_elements": {"layout": "title_only"},
                "speaker_notes": "Visual comparison emphasizes differences."
            },
            {
                "slide_number": 4,
                "slide_type": "content",
                "title": "Placeholder Example",
                "content_elements": [
                    {
                        "element_type": "text",
                        "bullet_points": ["Discuss future steps."]
                    },
                    {
                        "element_type": "image_placeholder",
                        "description": "Company Logo",
                        "position": {
                            "left_inches": 7.0, "top_inches": 5.0,
                            "width_inches": 2.0, "height_inches": 1.0
                        }
                    }
                ],
                "design_elements": {"layout": "title_content"},
                "speaker_notes": "Remember to insert logo later."
            },
            {
                "slide_number": 5,
                "slide_type": "closing",
                "title": "Conclusion & Contact",
                "content_elements": [
                    {
                        "element_type": "text",
                        "bullet_points": [
                            "AI testing is crucial for reliable systems.",
                            "Future work: Adversarial testing.",
                            "Contact: test@example.com"
                        ]
                    }
                ],
                "design_elements": {"layout": "title_content"},
                "speaker_notes": "End with a call to action or summary."
            }
        ],
        "template_info": sample_template_data
    }


@pytest.fixture
def mock_llm():
    """Provides a mock LLM that includes the bind_tools method."""
    mock = MagicMock(spec=["invoke", "bind", "bind_tools"])
    mock.bind_tools.return_value = mock

    return mock


@pytest.fixture
def mock_retriever():
    """Provides a mock Retriever."""
    mock = MagicMock(spec=["invoke", "get_relevant_documents"])
    mock.invoke.return_value = []
    return mock


@pytest.fixture
def mock_embeddings():
    """Provides a mock Embeddings object."""
    return MagicMock(spec=["embed_documents", "embed_query"])


@pytest.fixture
def mock_vector_store():
    """Provides a mock VectorStore."""
    mock = MagicMock(spec=["as_retriever", "similarity_search", "add_documents", "from_documents"])
    mock.from_documents.return_value = mock
    return mock


@pytest.fixture
def mock_agent_runnable():
    """Provides a mock agent runnable (like the output of create_react_agent)."""
    mock = MagicMock(spec=["invoke"])
    mock.invoke.return_value = {"output": json.dumps({"metadata": {}, "slides": [
        {"slide_number": 1, "title": "Mock Slide", "content_elements": [], "design_elements": {}}]})}
    return mock


@pytest.fixture
def mock_pptx_presentation(mock_pptx_slide):
    """ Mocks the pptx.Presentation object """
    mock_pres = MagicMock(name="Presentation")
    mock_pres.slides = MagicMock(name="Slides")
    mock_pres.slide_layouts = [
        MagicMock(name="Layout_0"), MagicMock(name="Layout_1"),
        MagicMock(name="Layout_2"), MagicMock(name="Layout_3"),
        MagicMock(name="Layout_4"), MagicMock(name="Layout_5"),
        MagicMock(name="Layout_6"), MagicMock(name="Layout_7"),
        MagicMock(name="Layout_8")
    ]

    mock_pres.slides.add_slide.return_value = mock_pptx_slide
    return mock_pres


@pytest.fixture
def mock_pptx_slide():
    """ Mocks a pptx.slide.Slide object """
    mock_slide = MagicMock(name="Slide")
    mock_slide.shapes = MagicMock(name="Shapes")
    mock_slide.placeholders = MagicMock(name="Placeholders")
    mock_slide.notes_slide = MagicMock(name="NotesSlide")
    mock_slide.notes_slide.notes_text_frame = MagicMock(name="NotesTextFrame")
    mock_slide.notes_slide.notes_text_frame.paragraphs = []
    mock_slide.notes_slide.notes_text_frame.add_paragraph.return_value = MagicMock()

    mock_title_ph = MagicMock(name="TitlePlaceholder")
    mock_title_ph.has_text_frame = True
    mock_title_ph.text_frame = MagicMock(name="TitleTextFrame")
    mock_title_para = MagicMock(name="TitleParagraph", level=0)
    mock_title_para.font = MagicMock(size=Pt(10))
    mock_title_para.runs = []

    mock_run_for_title = MagicMock(name="TitleRun",
                                   font=MagicMock(
                                       size=Pt(32),
                                       color=MagicMock(rgb=None)
                                   ))

    mock_title_para.add_run.return_value = mock_run_for_title
    mock_title_ph.text_frame.paragraphs = [mock_title_para]
    mock_slide.shapes.title = mock_title_ph

    mock_body_ph = MagicMock(name="BodyPlaceholder")
    mock_body_ph.has_text_frame = True
    mock_body_ph.text_frame = MagicMock(name="BodyTextFrame")
    mock_body_ph.text_frame.paragraphs = []

    mock_para_for_body = MagicMock(name="BodyParagraph",
                                   level=1,
                                   font=MagicMock(color=MagicMock(rgb=None)),
                                   text="")
    mock_body_ph.text_frame.add_paragraph.return_value = mock_para_for_body
    mock_body_ph.is_content_placeholder = True
    mock_body_ph.placeholder_format.idx = 1
    mock_slide.placeholders.__len__.return_value = 2
    mock_slide.placeholders.__getitem__.side_effect = lambda \
        idx: mock_body_ph if idx == 1 else mock_title_ph if idx == 0 else MagicMock()

    mock_table = MagicMock(name="Table")
    mock_table.columns = MagicMock()
    mock_column = MagicMock(width=None)
    mock_table.columns.__getitem__.return_value = mock_column
    mock_cell_text_frame = MagicMock(paragraphs=[MagicMock(font=MagicMock(color=MagicMock()))])
    mock_cell = MagicMock(text_frame=mock_cell_text_frame)
    mock_table.cell.return_value = mock_cell
    mock_add_table_return = MagicMock(table=mock_table)
    mock_slide.shapes.add_table.return_value = mock_add_table_return

    mock_textbox_tf = MagicMock(name="TextBoxTextFrame")
    mock_textbox_tf.paragraphs = []

    mock_para_for_textbox = MagicMock(name="TextBoxParagraph")
    mock_para_for_textbox.runs = []
    mock_run_for_textbox = MagicMock(name="TextBoxRun", text="", font=MagicMock(size=None, color=MagicMock()))

    def add_run_side_effect_textbox(*args, **kwargs):
        mock_para_for_textbox.runs.append(mock_run_for_textbox)
        return mock_run_for_textbox

    mock_para_for_textbox.add_run.side_effect = add_run_side_effect_textbox
    mock_para_for_textbox.clear = MagicMock(side_effect=lambda: setattr(mock_para_for_textbox, 'runs', []))

    def add_para_side_effect_textbox(*args, **kwargs):
        mock_para_for_textbox.runs = []
        mock_para_for_textbox.clear.reset_mock()
        mock_para_for_textbox.add_run.reset_mock()
        mock_textbox_tf.paragraphs.append(mock_para_for_textbox)
        return mock_para_for_textbox

    mock_textbox_tf.add_paragraph.side_effect = add_para_side_effect_textbox

    mock_slide.shapes.add_textbox.return_value = MagicMock(text_frame=mock_textbox_tf)
    mock_slide._mock_textbox_tf = mock_textbox_tf
    mock_slide.shapes.add_chart.return_value = MagicMock(
        chart=MagicMock(chart_title=MagicMock(text_frame=MagicMock(paragraphs=[MagicMock()])),
                        legend=MagicMock(),
                        plots=[MagicMock(has_data_labels=False,
                                         data_labels=MagicMock(font=MagicMock(color=MagicMock())))]))

    mock_slide.shapes.add_shape.return_value = MagicMock(
        name="MockedRectangle",
        text_frame=MagicMock(paragraphs=[MagicMock(font=MagicMock(color=MagicMock()))]),
        line=MagicMock(color=MagicMock(), dash_style=None),
        fill=MagicMock())

    return mock_slide
